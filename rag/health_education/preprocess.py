"""Cleans data/ahd_raw/AHD.xlsx into rag/health_education/data/ahd_cleaned.jsonl.

Run once (or whenever the raw file changes) via `python -m
rag.health_education.preprocess`. The raw file is never modified — this
only reads it (streaming, read_only=True, per data/AHD_DATA_INSPECTION.md's
same technique) and writes new files under rag/health_education/data/.

Per docs/AHD_DATA_PROVENANCE.md, this is a best-effort v1 cleanup, not a
clinical review: PII masking covers the one disclosure pattern actually
observed in the corpus ("اسمي X"), not a general PII scrubber; medication
content is over-inclusively quarantined (keyword hit, not a semantic
check) per the "bias toward the safer exclusion" principle already used
elsewhere in this project (rules/crisis.py, rules/red_flags.py).
"""

from __future__ import annotations

import json
import re
from dataclasses import dataclass
from pathlib import Path
from typing import Iterator

import openpyxl

_ARABIC_RE = re.compile(r"[؀-ۿ]")
_NAME_DISCLOSURE_RE = re.compile(r"اسمي\s+[؀-ۿ]+(?:\s+[؀-ۿ]+)?")
_JUNK_TEST_RE = re.compile(r"test\s*call", re.IGNORECASE)
_SCRAPED_LABEL_ANSWER_RE = re.compile(r"^أسئلة\s*وأجوبة\s*طبية")

# Same list used for data/AHD_DATA_INSPECTION.md's medication-content scan.
# Deliberately over-inclusive (a false-positive quarantine costs nothing;
# a missed dosage row reaching the index is the failure mode that matters —
# same asymmetry rules/crisis.py's own module docstring argues for).
MEDICATION_KEYWORDS: frozenset[str] = frozenset({
    "جرعة", "مجم", "ملغ", "مج/", "قرص", "أقراص", "حقنة", "حبة",
    "مضاد حيوي", "دواء ", "الجرعة اليومية", "مل/كغ", "mg", "ml",
})

_MIN_QUESTION_LEN = 5
_MAX_STORED_ANSWER_LEN = 2000

_DATA_DIR = Path(__file__).parent / "data"
_RAW_PATH = Path(__file__).resolve().parents[2] / "data" / "ahd_raw" / "AHD.xlsx"
_CLEANED_PATH = _DATA_DIR / "ahd_cleaned.jsonl"
_QUARANTINED_PATH = _DATA_DIR / "ahd_quarantined_medication.jsonl"

_SOURCE_METADATA = {
    "dataset": "AHD: Arabic Healthcare Dataset",
    "version": 5,
    "doi": "10.17632/mgj29ndgrk.5",
    "license": "CC BY 4.0",
    "source_url": "https://data.mendeley.com/datasets/mgj29ndgrk/5",
}


@dataclass(frozen=True)
class CleanResult:
    """Outcome of cleaning one raw (question, answer, category) row.

    Exactly one of `kept`, `quarantined_medication`, or neither (dropped
    as malformed/junk/non-Arabic) is true — this is what
    tests/unit/test_health_education_preprocess.py checks per case.
    """

    kept: bool
    quarantined_medication: bool
    pii_masked: bool
    dropped_reason: str | None  # "malformed" | "duplicate" | "junk" | "non_arabic" | None
    question: str
    answer: str


def mask_pii(text: str) -> tuple[str, bool]:
    """Replace a self-disclosed first name ("اسمي محمد") with a placeholder.

    Best-effort, not a general PII scrubber — see module docstring. Only
    the pattern actually confirmed present in this corpus
    (data/AHD_DATA_INSPECTION.md #5) is handled.
    """
    masked, count = _NAME_DISCLOSURE_RE.subn("اسمي [محجوب]", text)
    return masked, count > 0


def _is_medication_content(question: str, answer: str) -> bool:
    combined = f"{question} {answer}"
    return any(keyword in combined for keyword in MEDICATION_KEYWORDS)


def _is_junk(question: str, answer: str) -> bool:
    if len(question.strip()) < _MIN_QUESTION_LEN:
        return True
    if _JUNK_TEST_RE.search(question):
        return True
    if _SCRAPED_LABEL_ANSWER_RE.match(answer.strip()):
        return True
    return False


def clean_record(question: str | None, answer: str | None, category: str | None) -> CleanResult:
    """Pure per-row cleaning decision — no I/O. Used by both run() below
    and tests/unit/test_health_education_preprocess.py directly."""
    q = (question or "").strip()
    a = (answer or "").strip()
    c = (category or "").strip()

    if not q or not a or not c:
        return CleanResult(False, False, False, "malformed", q, a)

    if _is_junk(q, a):
        return CleanResult(False, False, False, "junk", q, a)

    if not _ARABIC_RE.search(q):
        return CleanResult(False, False, False, "non_arabic", q, a)

    q_masked, q_pii = mask_pii(q)
    a_masked, a_pii = mask_pii(a)
    pii_masked = q_pii or a_pii

    if _is_medication_content(q_masked, a_masked):
        return CleanResult(False, True, pii_masked, None, q_masked, a_masked)

    if len(a_masked) > _MAX_STORED_ANSWER_LEN:
        a_masked = a_masked[:_MAX_STORED_ANSWER_LEN].rstrip() + "…"

    return CleanResult(True, False, pii_masked, None, q_masked, a_masked)


def _iter_raw_rows(path: Path) -> Iterator[tuple[str | None, str | None, str | None]]:
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        worksheet = workbook[workbook.sheetnames[0]]
        for row in worksheet.iter_rows(min_row=2, values_only=True):
            question, answer, category = (row + (None, None, None))[:3]
            yield question, answer, category
    finally:
        workbook.close()


def run(raw_path: Path = _RAW_PATH, out_dir: Path = _DATA_DIR) -> dict[str, int]:
    """Stream-clean the raw AHD file into ahd_cleaned.jsonl +
    ahd_quarantined_medication.jsonl under out_dir. Returns counts for the
    provenance/report doc. The raw file itself is never opened for
    writing."""
    out_dir.mkdir(parents=True, exist_ok=True)
    cleaned_path = out_dir / "ahd_cleaned.jsonl"
    quarantined_path = out_dir / "ahd_quarantined_medication.jsonl"

    counts = {
        "raw_rows": 0,
        "malformed": 0,
        "duplicate": 0,
        "junk": 0,
        "non_arabic_excluded": 0,
        "pii_masked": 0,
        "quarantined_medication": 0,
        "retained": 0,
    }
    seen: set[tuple[str, str]] = set()
    record_id = 0

    with cleaned_path.open("w", encoding="utf-8") as cleaned_file, \
         quarantined_path.open("w", encoding="utf-8") as quarantined_file:
        for question, answer, category in _iter_raw_rows(raw_path):
            counts["raw_rows"] += 1
            result = clean_record(question, answer, category)

            if result.dropped_reason == "malformed":
                counts["malformed"] += 1
                continue
            if result.dropped_reason == "junk":
                counts["junk"] += 1
                continue
            if result.dropped_reason == "non_arabic":
                counts["non_arabic_excluded"] += 1
                continue

            dedupe_key = (result.question, result.answer)
            if dedupe_key in seen:
                counts["duplicate"] += 1
                continue
            seen.add(dedupe_key)

            if result.pii_masked:
                counts["pii_masked"] += 1

            record = {
                "id": record_id,
                "question": result.question,
                "answer": result.answer,
                "category": (category or "").strip(),
                "source": _SOURCE_METADATA,
            }
            record_id += 1

            if result.quarantined_medication:
                counts["quarantined_medication"] += 1
                quarantined_file.write(json.dumps(record, ensure_ascii=False) + "\n")
            else:
                counts["retained"] += 1
                cleaned_file.write(json.dumps(record, ensure_ascii=False) + "\n")

    return counts


if __name__ == "__main__":
    result_counts = run()
    print(json.dumps(result_counts, ensure_ascii=False, indent=2))
